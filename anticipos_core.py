"""Generador PDF 'SOLICITUD DE ANTICIPO DE CONTABILIDAD' (Formato Autorizacion Descuento).

Rellena el PDF template BV con datos del trabajador + monto + fecha usando
overlay (reportlab canvas encima del PDF original).

Coordenadas calibradas con pdfplumber sobre el template BV oficial.
"""
import io, os, re
from datetime import date

from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import white, black

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates')
TEMPLATE_ANTICIPO = os.path.join(TEMPLATES_DIR, 'anticipo_bv.pdf')

_MESES = ['','ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO','SETIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE']


def _reorder_nombre(nc):
    """NocoDB 'APELLIDO1 APELLIDO2 NOMBRE1 NOMBRE2' -> 'NOMBRE1 NOMBRE2 APELLIDO1 APELLIDO2'."""
    parts = (nc or '').split()
    if len(parts) >= 4:
        return ' '.join(parts[2:] + parts[:2])
    return nc or ''


def generate_anticipo_pdf(nombre_completo, dni, cargo, division, monto, fecha_solicitud=None):
    """Genera 1 PDF de solicitud de anticipo rellenado.

    Args:
      nombre_completo: str (viene ya en formato Nombres+Apellidos o Apellidos+Nombres,
        se reordena automaticamente si tiene 4+ tokens).
      dni: str
      cargo: str (puesto del trabajador)
      division: str (area/division, default 'Industria')
      monto: float o str (S/)
      fecha_solicitud: date | datetime | 'YYYY-MM-DD' | None (default hoy).
    Returns:
      bytes del PDF.
    """
    if fecha_solicitud is None:
        fecha_solicitud = date.today()
    elif isinstance(fecha_solicitud, str):
        y, m, d = fecha_solicitud[:10].split('-')
        fecha_solicitud = date(int(y), int(m), int(d))

    nombre_display = _reorder_nombre(nombre_completo)
    monto_str = f'{float(monto):,.2f}' if not isinstance(monto, str) else monto

    overlay_buf = io.BytesIO()
    c = canvas.Canvas(overlay_buf, pagesize=letter)  # US Letter 612x792
    c.setFont('Helvetica', 10)

    # Coordenadas calibradas: y_reportlab = 792 - top_pdfplumber + 7
    def y(top_pdfplumber):
        return 792 - top_pdfplumber + 7

    X = 205  # despues del label, sobre la linea de puntos
    c.drawString(X, y(198.3), nombre_display)
    c.drawString(X, y(232.9), str(dni or ''))
    c.drawString(X, y(267.3), str(cargo or ''))
    c.drawString(X, y(301.9), str(division or 'Industria'))
    c.drawString(X, y(388.2), monto_str)

    # Fecha: DD (x213) / MM (x262) / 202X - tapa el '202' del template y escribe año completo
    y_fecha = y(422.6)
    c.drawString(213, y_fecha, f'{fecha_solicitud.day:02d}')
    c.drawString(262, y_fecha, f'{fecha_solicitud.month:02d}')
    c.setFillColor(white)
    c.rect(287, y_fecha - 2, 22, 12, fill=1, stroke=0)
    c.setFillColor(black)
    c.drawString(289, y_fecha, str(fecha_solicitud.year))

    c.showPage()
    c.showPage()  # pagina 2 vacia (firma trabajador)
    c.save()
    overlay_buf.seek(0)

    template = PdfReader(TEMPLATE_ANTICIPO)
    overlay = PdfReader(overlay_buf)
    writer = PdfWriter()
    for i, page in enumerate(template.pages):
        if i < len(overlay.pages):
            page.merge_page(overlay.pages[i])
        writer.add_page(page)

    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


CAT_LABELS = [
    ('alimentacion_hospedaje', 'A. Alimentación y Hospedaje'),
    ('alquiler_equipos',       'B. Alquiler de Equipos'),
    ('transporte_movilizacion','C. Transporte y Movilización'),
    ('lavado_limpieza',        'D. Lavandería, Lavado y Cochera'),
    ('otros_bonos',            'E. Otros, Bonos y Copias'),
]


def generate_detalle_individual_xlsx(nombre, dni, cargo, cliente, mes_label, items, comments=None,
                                     columnas_originales=None, sheet_titulo=None):
    """Genera el xlsx detalle del trabajador.

    Modo 1 (nuevo, con `columnas_originales`): replica el formato del cuadro original
      (super-header, factores, headers, fila del trabajador).
    Modo 2 (fallback): tabla generica Categoria/Monto/Comentarios (compatibilidad).
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin = Side(border_style='thin', color='888888')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if columnas_originales:
        # ---- MODO CUADRO ORIGINAL ----
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Detalle Viatico'
        title_fill = PatternFill('solid', fgColor='F4B084')
        sub_fill = PatternFill('solid', fgColor='FCE4D6')
        title_font = Font(bold=True, size=12)
        hdr_font = Font(bold=True, size=10)

        # Row 1: super-header con titulo
        titulo = sheet_titulo or f'VIATICOS DEL MES {(mes_label or "").upper()}'
        # Determinar rango: cols 1..N segun columnas
        max_col = max((c['col_letter'] and openpyxl.utils.column_index_from_string(c['col_letter']) or 1)
                      for c in columnas_originales) if columnas_originales else 1
        max_col = max(max_col, 2)
        # Titulo abarca cols 2..max_col
        ws.cell(row=1, column=1, value='N°').fill = title_fill
        ws.cell(row=1, column=1).font = title_font
        ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row=1, column=1).border = border
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=max_col)
        c = ws.cell(row=1, column=2, value=titulo)
        c.fill = title_fill; c.font = title_font
        c.alignment = Alignment(horizontal='center', vertical='center'); c.border = border
        ws.row_dimensions[1].height = 22

        # Row 2: super-headers (IGNIFUGA, OFICINA) + factores
        has_super = any(co.get('header_super') for co in columnas_originales)
        has_factor = any(co.get('factor') is not None for co in columnas_originales)
        row_super = None; row_factor = None
        nxt = 2
        if has_super or has_factor:
            row_super = nxt; nxt += 1
        # (usamos misma fila para super y factor si solo hay una; pero mejor: super arriba, factor abajo)
        # Simplificacion: si hay ambos, usar 2 filas separadas
        row_factor = nxt if has_factor else None
        if has_factor: nxt += 1
        # Si solo hay super pero no factor, no incrementamos otra vez
        header_row_ws = nxt
        data_row_ws = nxt + 1

        if row_super and has_super:
            for co in columnas_originales:
                ci = openpyxl.utils.column_index_from_string(co['col_letter'])
                v = co.get('header_super') or ''
                if v:
                    c = ws.cell(row=row_super, column=ci, value=v)
                    c.fill = sub_fill; c.font = hdr_font
                    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    c.border = border
        if row_factor and has_factor:
            for co in columnas_originales:
                ci = openpyxl.utils.column_index_from_string(co['col_letter'])
                f_ = co.get('factor')
                if f_ is not None:
                    c = ws.cell(row=row_factor, column=ci, value=f_)
                    c.fill = sub_fill; c.font = Font(bold=True, size=9)
                    c.alignment = Alignment(horizontal='center', vertical='center')
                    c.border = border

        # Headers
        # col 1 = N°, col 2 = NOMBRE (implicito), luego cada columna_original
        c = ws.cell(row=header_row_ws, column=1, value='N°')
        c.fill = sub_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal='center'); c.border = border
        for co in columnas_originales:
            ci = openpyxl.utils.column_index_from_string(co['col_letter'])
            c = ws.cell(row=header_row_ws, column=ci, value=co['header'])
            c.fill = sub_fill; c.font = hdr_font
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = border
        ws.row_dimensions[header_row_ws].height = 34

        # Fila datos: N=1, nombre en col 2 si columnas_originales[0] tiene col_letter='B'
        c = ws.cell(row=data_row_ws, column=1, value=1)
        c.alignment = Alignment(horizontal='center'); c.border = border
        for co in columnas_originales:
            ci = openpyxl.utils.column_index_from_string(co['col_letter'])
            val = co.get('valor')
            # Si es la columna del nombre (primer col con nombre), poner el nombre real
            if 'PROYECTO' in (co['header'] or '').upper() or ('NOMBRE' in (co['header'] or '').upper()):
                val = nombre
            c = ws.cell(row=data_row_ws, column=ci, value=val)
            c.border = border
            if isinstance(val, (int, float)):
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal='right')
            else:
                c.alignment = Alignment(horizontal='center', wrap_text=True)

        # Widths por defecto
        default_widths = {1:5, 2:32}
        for co in columnas_originales:
            ci = openpyxl.utils.column_index_from_string(co['col_letter'])
            if ci not in default_widths:
                # 10 por defecto, o mas si header largo
                default_widths[ci] = max(10, min(20, len(co['header'] or '')//2 + 8))
        for ci, w_ in default_widths.items():
            ws.column_dimensions[get_column_letter(ci)].width = w_

        import io
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    # ---- MODO FALLBACK (categorias A-E) ----
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Detalle Viatico'

    thin = Side(border_style='thin', color='888888')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfill = PatternFill('solid', fgColor='1F4E78')
    hfont = Font(color='FFFFFF', bold=True, size=11)
    lbl = Font(bold=True, size=10)
    total_font = Font(bold=True, size=12, color='FFFFFF')
    total_fill = PatternFill('solid', fgColor='2E7D32')

    # Titulo
    ws.merge_cells('A1:C1')
    ws['A1'] = f'DETALLE DE VIATICO - {(mes_label or "").upper()}'
    ws['A1'].font = Font(bold=True, size=14, color='1F4E78')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    # Datos trabajador
    ws['A3'] = 'Trabajador:'; ws['A3'].font = lbl; ws['B3'] = nombre or ''
    ws.merge_cells('B3:C3')
    ws['A4'] = 'DNI/CE:';     ws['A4'].font = lbl; ws['B4'] = str(dni or '')
    ws.merge_cells('B4:C4')
    ws['A5'] = 'Cargo:';      ws['A5'].font = lbl; ws['B5'] = cargo or ''
    ws.merge_cells('B5:C5')
    ws['A6'] = 'Proyecto:';   ws['A6'].font = lbl; ws['B6'] = cliente or ''
    ws.merge_cells('B6:C6')

    # Tabla headers en fila 8
    headers = ['Categoria', 'Monto (S/)', 'Comentarios']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=8, column=i, value=h)
        c.fill = hfill; c.font = hfont
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
    ws.row_dimensions[8].height = 26

    # Filas de categorias
    items = items or {}
    comments = comments or {}
    total = 0.0
    r = 9
    for key, label in CAT_LABELS:
        val = float(items.get(key, 0) or 0)
        cmt = comments.get(key, '') or ''
        ws.cell(row=r, column=1, value=label).border = border
        cm = ws.cell(row=r, column=2, value=val); cm.number_format = '#,##0.00'
        cm.border = border; cm.alignment = Alignment(horizontal='right')
        cc = ws.cell(row=r, column=3, value=cmt); cc.border = border
        cc.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        total += val
        r += 1

    # Total
    ws.cell(row=r, column=1, value='TOTAL').font = total_font
    ws.cell(row=r, column=1).fill = total_fill
    ws.cell(row=r, column=1).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=r, column=1).border = border
    tc = ws.cell(row=r, column=2, value=total)
    tc.number_format = '#,##0.00'; tc.font = total_font; tc.fill = total_fill
    tc.alignment = Alignment(horizontal='right', vertical='center'); tc.border = border
    ws.cell(row=r, column=3, value='').fill = total_fill
    ws.cell(row=r, column=3).border = border
    ws.row_dimensions[r].height = 22

    # Widths
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 45

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_anticipos_batch(workers_with_personal, fecha_solicitud=None):
    """Genera 1 PDF por trabajador matcheado (con monto > 0).

    workers_with_personal: [{'worker': {..., 'total'}, 'personal': {..., 'dni', 'nombre_completo', 'puesto', 'division'}}, ...]

    Returns:
      list de dicts {'filename', 'pdf_bytes', 'dni', 'nombre', 'monto'}.
    """
    out = []
    for item in workers_with_personal:
        p = item.get('personal') or {}
        w = item.get('worker') or {}
        if not p: continue
        monto = float(w.get('total', 0) or 0)
        if monto <= 0: continue
        dni = str(p.get('dni', ''))
        nombre = p.get('nombre_completo', '') or ''
        pdf = generate_anticipo_pdf(
            nombre_completo=nombre,
            dni=dni,
            cargo=p.get('puesto', '') or w.get('cargo', ''),
            division=p.get('division', '') or 'Industria',
            monto=monto,
            fecha_solicitud=fecha_solicitud,
        )
        safe = re.sub(r'[^A-Za-z0-9_-]', '_', nombre)[:40]
        out.append({
            'filename': f'DJ_Anticipo_{dni}_{safe}.pdf',
            'pdf_bytes': pdf,
            'dni': dni,
            'nombre': nombre,
            'monto': monto,
        })
    return out
